import csv
import io
from datetime import UTC, date, datetime, timedelta

import pytest
from bson import ObjectId
from openpyxl import load_workbook

from app.domain.exports import ExportFormat, export_filename, neutralize_formula
from app.domain.permissions import Role
from app.exports.pdf_writer import pdf_text


@pytest.fixture
def people(make_user):
    return {
        "lead": make_user("Lead", Role.TECH_LEAD),
        "dev": make_user("Dev", Role.DEVELOPER),
        "outsider": make_user("Outsider", Role.TECH_LEAD),
    }


@pytest.fixture
def apis(login_as, people):
    return {name: login_as(user) for name, user in people.items()}


@pytest.fixture
def project(apis, people):
    response = apis["lead"].post(
        "/api/projects",
        json={"key": "EXP", "name": "Exportações", "memberIds": [str(people["dev"].id)]},
    )
    assert response.status_code == 201, response.text
    return response.json()


@pytest.fixture
def issues(apis, people, project, indexed_db):
    lead = apis["lead"]

    def create(**overrides):
        body = {"projectId": project["id"], "type": "story", "title": "Funcionalidade"}
        response = lead.post("/api/issues", json=body | overrides)
        assert response.status_code == 201, response.text
        return response.json()

    def patch(issue, **changes):
        response = lead.patch(
            f"/api/issues/{issue['id']}", json={"version": issue["version"]} | changes
        )
        assert response.status_code == 200, response.text
        return response.json()

    first = create(title="Login", assigneeId=str(people["dev"].id), storyPoints=3)
    patch(first, status="in_progress")
    create(title='=HYPERLINK("http://x")', type="bug", priority="critical")
    old = create(title="Antiga")
    old = patch(patch(old, status="in_progress"), status="done")
    indexed_db["issues"].update_one(
        {"_id": ObjectId(old["id"])},
        {"$set": {"completedAt": datetime.now(UTC) - timedelta(days=200)}},
    )
    done = create(title="Feita — com travessão")
    patch(patch(done, status="in_progress"), status="done")


def _export(api, project, export_format, **params):
    query = {"projectId": project["id"], "format": export_format} | params
    return api.get("/api/reports/export", params=query)


def test_domain_helpers():
    assert neutralize_formula("=1+1") == "'=1+1"
    assert neutralize_formula("@SUM(A1)") == "'@SUM(A1)"
    assert neutralize_formula("-2") == "'-2"
    assert neutralize_formula("Login") == "Login"
    assert (
        export_filename("A/B", ExportFormat.PDF, date(2026, 1, 2))
        == "vccteamforge-AB-metricas-2026-01-02.pdf"
    )
    assert pdf_text("Feita — já 🚀") == "Feita - já ?"


def test_csv_export(apis, project, issues, indexed_db):
    response = _export(apis["dev"], project, "csv")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    today = datetime.now(UTC).date().isoformat()
    assert response.headers["content-disposition"] == (
        f'attachment; filename="vccteamforge-EXP-metricas-{today}.csv"'
    )
    assert response.headers["cache-control"] == "no-store"
    assert response.content.startswith(b"\xef\xbb\xbf")

    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][:6] == ["Chave", "Título", "Tipo", "Estado", "Prioridade", "Responsável"]
    # A issue concluída antes do período fica de fora; ordem pelo número da issue.
    assert [row[0] for row in rows[1:]] == ["EXP-1", "EXP-2", "EXP-4"]
    assert rows[1][3:8] == ["Em progresso", "Média", "Dev", "", "3"]
    assert rows[2][1] == '\'=HYPERLINK("http://x")'
    assert rows[2][5] == "Sem responsável"

    audit = indexed_db["audit_logs"].find_one({"action": "report.exported"})
    assert audit["details"]["format"] == "csv"
    assert audit["details"]["issueCount"] == 3
    assert audit["details"]["filters"]["type"] is None

    bugs = _export(apis["dev"], project, "csv", type="bug")
    assert [row[0] for row in csv.reader(io.StringIO(bugs.text), delimiter=";")][1:] == ["EXP-2"]


def test_xlsx_export(apis, project, issues):
    response = _export(apis["lead"], project, "xlsx")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == [
        "Resumo",
        "Issues",
        "Estados",
        "Fluxo",
        "Velocity",
        "Aging",
        "Mais antigas",
        "Bugs por severidade",
        "QA",
        "Carga de trabalho",
        "Releases",
    ]
    summary = {row[0]: row[1] for row in workbook["Resumo"].iter_rows(values_only=True) if row[0]}
    assert summary["Projeto"] == "EXP — Exportações"
    assert summary["Abertas"] == 2
    assert summary["Bugs críticos em aberto"] == 1

    sheet = workbook["Issues"]
    formula = sheet["B3"]
    assert formula.value == '=HYPERLINK("http://x")'
    assert formula.data_type == "s"
    assert isinstance(sheet["J2"].value, datetime)
    states = {row[0]: row[1] for row in workbook["Estados"].iter_rows(min_row=2, values_only=True)}
    assert states["Concluída"] == 2


def test_pdf_export(apis, project, issues):
    response = _export(apis["dev"], project, "pdf", type="story")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].endswith('.pdf"')
    assert response.content.startswith(b"%PDF-")
    assert len(response.content) > 1000


def test_export_validation_and_access(apis, project, issues, monkeypatch):
    dev = apis["dev"]
    assert _export(apis["outsider"], project, "csv").status_code == 404
    assert dev.get("/api/reports/export", params={"projectId": project["id"]}).status_code == 422
    assert _export(dev, project, "docx").status_code == 422

    period = _export(dev, project, "csv", **{"from": "2026-02-01", "to": "2026-01-01"})
    assert period.status_code == 422
    assert period.json()["field"] == "from"

    other = apis["outsider"].post("/api/projects", json={"key": "OUT", "name": "Alheio"}).json()
    sprint = (
        apis["outsider"]
        .post(
            "/api/sprints",
            json={
                "projectId": other["id"],
                "name": "Alheio",
                "startDate": date.today().isoformat(),
                "endDate": (date.today() + timedelta(days=9)).isoformat(),
            },
        )
        .json()
    )
    foreign = _export(dev, project, "csv", sprintId=sprint["id"])
    assert foreign.status_code == 422
    assert foreign.json()["field"] == "sprintId"

    monkeypatch.setattr("app.services.report_exports.EXPORT_MAX_ISSUES", 2)
    limited = _export(dev, project, "xlsx")
    assert limited.status_code == 422
    assert "limitada a 2 issues (3 encontradas)" in limited.json()["detail"]
